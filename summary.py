"""Client summary (template or optional Azure OpenAI wording) and Word/Excel exports."""
import io
import json
import re

import pandas as pd

from compare import CHANGE_TYPES
from impact import FLAGS
from ui_format import fmt_int, fmt_money, fmt_pct, fmt_signed_money

INK = "17324D"
FOOTER = "All figures computed from public fee schedule data and synthetic claims."


def _label(code, mod):
    mod = "" if mod is None or (isinstance(mod, float) and pd.isna(mod)) else str(mod)
    return f"{code}-{mod}" if mod else str(code)


def md_safe(text):
    """Escape $ so Streamlit Markdown never renders it as LaTeX."""
    return re.sub(r"(?<!\\)\$", r"\\$", text)


# ----------------------------------------------------------------------------- stats
def build_stats(meta, changes_summary, impact, config):
    top = impact["top_10"].head(5)
    top_changes = [{
        "code": _label(r.code, r.modifier),
        "change": r.change_type,
        "old_rate": fmt_money(r.old_rate, 2),
        "new_rate": fmt_money(r.new_rate, 2),
        "impact": fmt_signed_money(r.impact),
    } for r in top.itertuples()]

    flags = config["flags_df"]
    out = flags[flags["flag"] == FLAGS[0]].copy()
    out["_abs"] = out["payment_error"].abs()
    out = out.sort_values("_abs", ascending=False).head(5)
    cpf = config["count_per_flag"]
    net = impact["net_impact"]
    return {
        "dataset": meta.get("name", "Fee schedule"),
        "old_version": meta.get("old_label", "Old version"),
        "new_version": meta.get("new_label", "New version"),
        "source": meta.get("source", ""),
        "claims_note": meta.get("claims_note", "synthetic claim lines"),
        "pct_of_fee_schedule": f"{impact['pct_of_fee_schedule']:g}%",
        "codes_compared": fmt_int(changes_summary["total_rows"]),
        "counts": {t: int(changes_summary.get(t, 0)) for t in CHANGE_TYPES},
        "net_impact": fmt_money(abs(net)),
        "net_impact_signed": fmt_signed_money(net),
        "net_direction": "higher cost" if net >= 0 else "savings",
        "total_increase": fmt_money(impact["total_increase"]),
        "total_decrease": fmt_money(impact["total_decrease"]),
        "decrease_abs": fmt_money(abs(impact["total_decrease"])),
        "top10_share_of_increase": fmt_pct(impact["top10_share_of_increase"]),
        "top_changes": top_changes,
        "deleted_code_units": fmt_int(impact["deleted_code_volume"]),
        "config_issues": int(config["issues"]),
        "outdated": int(cpf[FLAGS[0]]),
        "not_configured": int(cpf[FLAGS[1]]),
        "retired": int(cpf[FLAGS[2]]),
        "not_priced_new": int(changes_summary.get("Became not priced", 0)),
        "overpayment": fmt_money(config["total_overpayment"]),
        "underpayment": fmt_money(config["total_underpayment"]),
        "underpayment_abs": fmt_money(abs(config["total_underpayment"])),
        "invalid_payments": fmt_money(config["invalid_code_payments"]),
        "total_mispaid": fmt_money(config["total_mispaid"]),
        "outdated_top_codes": ", ".join(_label(c, m) for c, m in zip(out["code"], out["modifier"])),
        "outdated_top_dollars": fmt_money(float(out["_abs"].sum())),
        "units_not_configured": fmt_int(config["units_not_configured"]),
    }


def recommendations(stats):
    recs = []
    if stats["outdated"]:
        recs.append(f"**Priority 1. Update {stats['outdated']} outdated rates.** The five largest "
                    f"({stats['outdated_top_codes']}) account for {stats['outdated_top_dollars']} "
                    f"of mispayment.")
    if stats["retired"]:
        recs.append(f"**Priority 2. Retire {stats['retired']} deleted codes still active in the "
                    f"system.** They paid {stats['invalid_payments']} on codes no longer in the "
                    f"schedule.")
    if stats["not_configured"]:
        recs.append(f"**Priority 3. Add {stats['not_configured']} new codes.** "
                    f"{stats['units_not_configured']} claim units will pend or deny until they "
                    f"are configured.")
    if stats["not_priced_new"]:
        recs.append(f"**Priority 4. Set manual pricing for {stats['not_priced_new']} codes that "
                    f"became not priced.**")
    recs.append("Load the generated configuration update file, then reprice a sample of claims "
                "to confirm.")
    return recs


# ----------------------------------------------------------------------------- template
def template_summary(stats):
    c = stats["counts"]
    tc = stats["top_changes"]
    drivers = "; ".join(f"{t['code']} ({t['change'].lower()}, {t['old_rate']} to "
                        f"{t['new_rate']}, {t['impact']})" for t in tc[:3]) or "none"
    if stats["net_direction"] == "higher cost":
        headline = f"a net {stats['net_impact']} higher annual cost"
    else:
        headline = f"net annual savings of {stats['net_impact']}"
    lines = [
        "### Overview",
        f"We compared {stats['old_version']} and {stats['new_version']} of the {stats['dataset']} "
        f"across {stats['codes_compared']} code lines. Against {stats['claims_note']}, the update "
        f"means {headline}, and {stats['config_issues']} configuration issues would cause "
        f"{stats['total_mispaid']} to be paid incorrectly.",
        "",
        "### Key rate changes",
        f"- {c['Rate increase']:,} rate increases and {c['Rate decrease']:,} rate decreases",
        f"- {c['New code']:,} new codes and {c['Deleted code']:,} deleted codes",
        f"- {c['Became not priced']:,} codes became not priced; {c['Became priced']:,} became priced",
        f"- {c['No change']:,} code lines kept the same rate",
        "",
        "### Financial impact",
        f"Increases add {stats['total_increase']} and decreases save {stats['decrease_abs']}, "
        f"for {headline}. The top 10 codes drive {stats['top10_share_of_increase']} of the "
        f"increase. The three largest drivers are {drivers}.",
        "",
        "### Configuration risks",
        f"- {stats['outdated']} outdated rates: {stats['overpayment']} overpaid and "
        f"{stats['underpayment_abs']} underpaid",
        f"- {stats['retired']} retired codes still active: {stats['invalid_payments']} paid on "
        f"invalid codes",
        f"- {stats['not_configured']} new codes not configured: {stats['units_not_configured']} "
        f"claim units will pend or deny",
        "",
        "### Recommended actions",
    ]
    lines += [f"- {r}" for r in recommendations(stats)]
    lines += [
        "",
        "### Assumptions",
        f"- Claims volume: {stats['claims_note']}, treated as 12 months",
        f"- The plan pays {stats['pct_of_fee_schedule']} of the state fee schedule",
        f"- Source: {stats['source']}; no client or member data",
    ]
    return "\n".join(lines)


SYSTEM_PROMPT = (
    "You are a healthcare payment accuracy consultant. Write a client-ready summary of a "
    "Medicaid fee schedule update using ONLY the figures in the JSON provided. Copy every figure "
    "exactly as written; do not calculate, round, estimate or invent any number. Use exactly "
    "these six Markdown sections with ### headings: Overview, Key rate changes, Financial impact, "
    "Configuration risks, Recommended actions, Assumptions. Keep the recommendations in the "
    "priority order given. Plain business English, under 320 words."
)


def ai_summary(stats, endpoint, key, deployment):
    from openai import AzureOpenAI  # optional dependency

    client = AzureOpenAI(azure_endpoint=endpoint, api_key=key, api_version="2024-06-01")
    payload = dict(stats)
    payload["recommendations"] = recommendations(stats)
    resp = client.chat.completions.create(
        model=deployment,
        temperature=0.2,
        messages=[{"role": "system", "content": SYSTEM_PROMPT},
                  {"role": "user", "content": json.dumps(payload)}],
    )
    text = (resp.choices[0].message.content or "").strip()
    if not text:
        raise RuntimeError("The AI returned an empty summary.")
    return text


def build_summary(stats, use_ai=False, secrets=None):
    """Return (text, source, error). Never raises."""
    error = None
    if use_ai:
        try:
            s = secrets or {}
            text = ai_summary(stats, s["AZURE_OPENAI_ENDPOINT"], s["AZURE_OPENAI_KEY"],
                              s["AZURE_OPENAI_DEPLOYMENT"])
            return text, "Azure OpenAI (wording only; figures from code)", None
        except Exception as exc:  # noqa: BLE001
            error = f"AI rewrite failed, so the template summary is shown. ({exc})"
    try:
        return template_summary(stats), "Template (built from computed figures)", error
    except Exception as exc:  # noqa: BLE001
        return "Summary could not be built.", "None", str(exc)


# ----------------------------------------------------------------------------- exports
def _add_runs(par, text):
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            run = par.add_run(part.replace("\\$", "$"))
            run.bold = i % 2 == 1


def export_docx(markdown, title, subtitle=""):
    from docx import Document
    from docx.shared import RGBColor

    doc = Document()
    h = doc.add_heading(title, level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor.from_string(INK)
    if subtitle:
        p = doc.add_paragraph()
        p.add_run(subtitle).italic = True
    for line in markdown.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("#"):
            hh = doc.add_heading(s.lstrip("#").strip(), level=2)
            for r in hh.runs:
                r.font.color.rgb = RGBColor.from_string(INK)
        elif s.startswith("- "):
            _add_runs(doc.add_paragraph(style="List Bullet"), s[2:])
        else:
            _add_runs(doc.add_paragraph(), s)
    p = doc.add_paragraph()
    p.add_run(FOOTER).italic = True
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def export_excel(sheets):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=INK)
            for i, col in enumerate(df.columns, start=1):
                lens = [len(str(col))] + [len(str(v)) for v in df[col].head(500).tolist()]
                ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, max(lens) + 2))
            ws.freeze_panes = "A2"
    buf.seek(0)
    return buf
